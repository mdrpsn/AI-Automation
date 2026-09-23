#!/usr/bin/env python3
"""
Builds a monitoring Excel sheet (roster x date) from the .txt files in
"Previous Schedules/", for a given date range. Each cell shows where/what
shift a person worked, or OFF / LEAVE.

Usage:
  python scheduler/generate_monitoring.py --start "08 01 2026" --end "08 08 2026"
"""
import argparse
import sys
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from generate_schedule import SCHED_DIR, load_roster, parse_file  # noqa: E402

WORK_FILL = PatternFill("solid", fgColor="C6E0B4")   # green
OFF_FILL = PatternFill("solid", fgColor="FFE699")    # yellow
LEAVE_FILL = PatternFill("solid", fgColor="F8CBAD")  # orange/red
BLANK_FILL = PatternFill("solid", fgColor="D9D9D9")  # grey (unaccounted)
HEADER_FILL = PatternFill("solid", fgColor="2F5496")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SECTION_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FONT = Font(color="FFFFFF", bold=True, size=12)

# Order and label for role-based sections. A person's roster.json "roles"
# entry ("Barista" / "Kitchen" / "Both") puts them in exactly one section.
ROLE_SECTIONS = [
    ("BARISTA", "Barista"),
    ("KITCHEN", "Kitchen"),
    ("BOTH (BARISTA OR KITCHEN)", "Both"),
]


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--start", required=True, help="Start date 'MM DD YYYY'")
    ap.add_argument("--end", required=True, help="End date 'MM DD YYYY'")
    ap.add_argument("--out", default=None, help="Output .xlsx path")
    args = ap.parse_args()

    start = datetime.strptime(args.start, "%m %d %Y").date()
    end = datetime.strptime(args.end, "%m %d %Y").date()

    roster_cfg = load_roster()
    aliases = roster_cfg.get("name_aliases", {})
    roster = roster_cfg["roster"]
    roles_map = roster_cfg.get("roles", {})

    dates = []
    d = start
    while d <= end:
        dates.append(d)
        d += timedelta(days=1)

    day_data = {}  # date -> (worked, off)
    missing = []
    for d in dates:
        path = SCHED_DIR / f"{d.strftime('%m %d %Y')}.txt"
        if not path.exists():
            missing.append(d)
            continue
        _, worked, off = parse_file(path, aliases)
        day_data[d] = (worked, off)

    leave_by_date = {}
    for d in dates:
        leaves = set()
        for name, value in roster_cfg.get("on_leave", {}).items():
            # A person's on_leave value is either a single {since,until,note}
            # dict or a list of them, for people with more than one leave
            # period on record -- same normalization as generate_schedule.py's
            # active_leave(), which this duplicates rather than imports.
            periods = value if isinstance(value, list) else [value]
            for info in periods:
                since = datetime.strptime(info["since"], "%Y-%m-%d").date()
                until = datetime.strptime(info["until"], "%Y-%m-%d").date() if info.get("until") else None
                if since <= d and (until is None or d <= until):
                    leaves.add(name)
                    break
        leave_by_date[d] = leaves

    wb = Workbook()
    ws = wb.active
    ws.title = "Monitoring"

    ws.cell(row=1, column=1, value="NAME").font = HEADER_FONT
    ws.cell(row=1, column=1).fill = HEADER_FILL
    for ci, d in enumerate(dates, start=2):
        c = ws.cell(row=1, column=ci, value=d.strftime("%b %d\n%a"))
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    summary_col = len(dates) + 2
    for label, offset in (("DAYS WORKED", 0), ("DAYS OFF", 1), ("DAYS LEAVE", 2)):
        c = ws.cell(row=1, column=summary_col + offset, value=label)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    grouped = {role: sorted(n for n in roster if roles_map.get(n) == role) for _, role in ROLE_SECTIONS}
    unclassified = sorted(n for n in roster if n not in roles_map)
    sections = [(label, names) for label, role in ROLE_SECTIONS if (names := grouped[role])]
    if unclassified:
        sections.append(("UNCLASSIFIED (add to roster.json roles)", unclassified))

    last_col = summary_col + 2
    ri = 2
    for label, names in sections:
        hc = ws.cell(row=ri, column=1, value=label)
        hc.font = SECTION_FONT
        for ci in range(1, last_col + 1):
            ws.cell(row=ri, column=ci).fill = SECTION_FILL
        ws.merge_cells(start_row=ri, start_column=1, end_row=ri, end_column=last_col)
        ri += 1

        for name in names:
            nc = ws.cell(row=ri, column=1, value=name)
            nc.font = Font(bold=True)

            worked_count = off_count = leave_count = 0
            for ci, d in enumerate(dates, start=2):
                cell = ws.cell(row=ri, column=ci)
                if d not in day_data:
                    cell.value = "?"
                    cell.fill = BLANK_FILL
                    cell.alignment = Alignment(horizontal="center")
                    continue

                worked, off = day_data[d]
                leave_today = leave_by_date[d]

                if name in leave_today:
                    cell.value = "LEAVE"
                    cell.fill = LEAVE_FILL
                    leave_count += 1
                elif name in worked:
                    lines = []
                    for loc, shift, role in worked[name]:
                        shift_part = f" {shift}" if shift and shift != "-" else ""
                        lines.append(f"{loc}{shift_part} {role}")
                    cell.value = "\n".join(lines)
                    cell.fill = WORK_FILL
                    worked_count += 1
                elif name in off:
                    cell.value = "OFF"
                    cell.fill = OFF_FILL
                    off_count += 1
                else:
                    cell.value = "-"
                    cell.fill = BLANK_FILL
                cell.alignment = Alignment(horizontal="center", wrap_text=True)

            ws.cell(row=ri, column=summary_col, value=worked_count).alignment = Alignment(horizontal="center")
            ws.cell(row=ri, column=summary_col + 1, value=off_count).alignment = Alignment(horizontal="center")
            ws.cell(row=ri, column=summary_col + 2, value=leave_count).alignment = Alignment(horizontal="center")
            ri += 1

    ws.column_dimensions["A"].width = 14
    for ci in range(2, len(dates) + 2):
        ws.column_dimensions[get_column_letter(ci)].width = 13
    for offset in range(3):
        ws.column_dimensions[get_column_letter(summary_col + offset)].width = 12
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 30

    out_path = Path(args.out) if args.out else Path(__file__).resolve().parent.parent / \
        f"Monitoring {start.strftime('%m-%d')} to {end.strftime('%m-%d %Y')}.xlsx"
    wb.save(out_path)
    print(f"Written to {out_path}")
    if missing:
        print("WARNING: no schedule file found for:", ", ".join(d.strftime("%m %d %Y") for d in missing))


if __name__ == "__main__":
    main()
